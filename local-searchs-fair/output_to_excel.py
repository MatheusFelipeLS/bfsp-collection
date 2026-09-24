#!/usr/bin/env python3
"""Agrega os results/<entrada>.csv gerados pelo run_all.py (uma linha bruta
por ITERAÇÃO, todas rodadas sob o MESMO orçamento de tempo - ver README.md
"Comparação justa") num único output.xlsx com os valores médios por
instância, incluindo o RPD (Relative Percentage Deviation).

Formato de results/<entrada>.csv (cabeçalho + `iters` linhas por instância,
uma por repetição - sem nenhuma agregação prévia):
    instance,iteration,initial_cost,final_cost,improvement_abs,improvement_pct,time_ms,time_limit_ms,note
    J20M5N1,1,1828,1419,409,22.3742,0.064466,2000,passes=2
    J20M5N1,2,1828,1419,409,22.3742,0.095958,2000,passes=2
    ...

Por que RPD em vez de "Melhoria (%)": Melhoria (%) mede o quanto cada busca
melhorou a MESMA solução inicial aleatória - não diz nada sobre quão perto do
melhor resultado possível ela chegou, então um método "fraco" numa instância
fácil pode parecer melhor que um método forte numa instância difícil. RPD
mede a distância até uma referência fixa por instância:

    RPD = (custo_médio - referência) / referência * 100  (0% = igualou a referência)

A referência é, por instância, o melhor valor conhecido da literatura
(best_ribas.json na raiz do repo, Taillard/Ribas) quando disponível; quando
não (hoje, só as instâncias J500M5/J500M10, que não estão nesse arquivo),
cai para o melhor custo médio obtido por QUALQUER entrada nesta própria
rodada - a coluna "Fonte RPD" de cada linha diz qual foi usada.

Este script agrupa as linhas brutas por instância e calcula, por entrada, uma
aba com uma linha por instância:

    Instância | Custo inicial | Custo final (méd) | RPD (%) | Fonte RPD
             | Tempo médio (ms) | Tempo mín (ms) | Tempo máx (ms)
             | Orçamento (ms) | Uso do orçamento (%) | N iterações | Nota

Uma aba "Resumo" no início lista, por entrada, o RPD médio/mín/máx e o uso
médio do orçamento - ordenada do menor (melhor) para o maior RPD médio, ou
seja, já funciona como um ranking.
"""

import argparse
import csv
import json
import math
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

# Ordem do catálogo (#1..#23) do README, usada nas abas.
ENTRY_ORDER = [
    "rls",
    "rls-grabowski",
    "rls-de-abc",
    "mrls-p-eda",
    "best-swap-ig",
    "best-swap-ig-fixed",
    "swap-first-ig",
    "vnd1",
    "vnd2",
    "vnd1-fixed",
    "vnd2-fixed",
    "ig-ij-dispatch",
    "svns-s-swap",
    "svns-s-insertion",
    "svns-d-swap",
    "svns-d-insertion",
    "hvns-best-insertion",
    "hvns-best-edge-insertion",
    "hvns-best-swap",
    "hvns-shaking",
    "hvns-sa-rls",
    "hvns-sa-edge-insertion",
    "tpa-sa",
]

SHEET_HEADER = [
    "Instância", "Custo inicial", "Custo final (méd)", "RPD (%)", "Fonte RPD", "Tempo médio (ms)", "Tempo mín (ms)",
    "Tempo máx (ms)", "Orçamento (ms)", "Uso do orçamento (%)", "N iterações", "Nota",
]

INSTANCE_RE = re.compile(r"^(J\d+M\d+)N(\d+)$")


def load_literature_best(path):
    """Retorna {"J20M5": [custo_N1, custo_N2, ...], ...} a partir de
    best_ribas.json (Taillard/Ribas), ou {} se o arquivo não existir."""
    if not os.path.exists(path):
        return {}
    with open(path) as f:
        return json.load(f)


def literature_best_cost(instance_name, literature):
    m = INSTANCE_RE.match(instance_name)
    if not m:
        return None
    group, index = m.group(1), int(m.group(2)) - 1
    values = literature.get(group)
    if values is None or index >= len(values):
        return None
    return values[index]


def parse_results_file(path):
    """Retorna [{instance, iteration, initial_cost, final_cost, improvement_abs,
    improvement_pct, time_ms, time_limit_ms, note}, ...] - uma entrada por
    iteração bruta."""
    rows = []
    with open(path, newline="") as f:
        reader = csv.DictReader(f)
        for rec in reader:
            rows.append({
                "instance": rec["instance"],
                "iteration": int(rec["iteration"]),
                "initial_cost": int(rec["initial_cost"]),
                "final_cost": int(rec["final_cost"]),
                "improvement_abs": float(rec["improvement_abs"]),
                "improvement_pct": float(rec["improvement_pct"]),
                "time_ms": float(rec["time_ms"]),
                "time_limit_ms": float(rec["time_limit_ms"]),
                "note": rec["note"],
            })
    return rows


def aggregate_by_instance(rows):
    """Agrupa as linhas brutas (uma por iteração) por instância e retorna uma
    linha agregada por instância (sem RPD ainda - isso depende de conhecer
    todas as entradas, feito depois em attach_rpd), na ordem de primeira
    aparição."""
    order = []
    by_instance = {}
    for r in rows:
        key = r["instance"]
        if key not in by_instance:
            order.append(key)
            by_instance[key] = []
        by_instance[key].append(r)

    aggregated = []
    for instance in order:
        group = by_instance[instance]
        times = [r["time_ms"] for r in group]
        time_limit_ms = group[0]["time_limit_ms"]
        time_ms_mean = sum(times) / len(times)
        aggregated.append({
            "instance": instance,
            "initial_cost": group[0]["initial_cost"],
            "final_cost_mean": sum(r["final_cost"] for r in group) / len(group),
            "improvement_abs_mean": sum(r["improvement_abs"] for r in group) / len(group),
            "improvement_pct_mean": sum(r["improvement_pct"] for r in group) / len(group),
            "time_ms_mean": time_ms_mean,
            "time_ms_min": min(times),
            "time_ms_max": max(times),
            "time_limit_ms": time_limit_ms,
            "budget_used_pct": (time_ms_mean / time_limit_ms * 100) if time_limit_ms > 0 else 0.0,
            "n_iterations": len(group),
            "note": group[-1]["note"],
        })
    return aggregated


def attach_rpd(aggregated_by_entry, literature):
    """Calcula RPD para cada linha agregada (mutando-as in-place): usa o
    melhor conhecido da literatura quando disponível, senão o melhor custo
    médio encontrado por QUALQUER entrada nesta rodada para aquela
    instância."""
    best_found = {}
    for agg_rows in aggregated_by_entry.values():
        for r in agg_rows:
            best_found[r["instance"]] = min(best_found.get(r["instance"], math.inf), r["final_cost_mean"])

    for agg_rows in aggregated_by_entry.values():
        for r in agg_rows:
            lit_best = literature_best_cost(r["instance"], literature)
            if lit_best:
                reference, source = lit_best, "literatura (Ribas)"
            else:
                reference, source = best_found[r["instance"]], "melhor desta rodada"
            r["rpd"] = ((r["final_cost_mean"] - reference) / reference * 100) if reference > 0 else 0.0
            r["rpd_source"] = source


def discover_results(results_dir):
    """Retorna {entrada: rows_brutas} (uma linha por iteração) para os
    results/<entrada>.csv existentes, na ordem de ENTRY_ORDER (extras vão
    para o fim, em ordem alfabética)."""
    found = {}
    for name in os.listdir(results_dir):
        if not name.endswith(".csv"):
            continue
        entry = os.path.splitext(name)[0]
        found[entry] = parse_results_file(os.path.join(results_dir, name))
    if not found:
        sys.exit(f"nenhum results/*.csv encontrado em {results_dir} (rode run_all.py antes)")

    ordered = [e for e in ENTRY_ORDER if e in found]
    ordered += sorted(e for e in found if e not in ENTRY_ORDER)
    return {e: found[e] for e in ordered}


def autosize_columns(sheet):
    """Ajusta a largura de cada coluna ao maior valor (cabeçalho incluído)."""
    for col in sheet.iter_cols():
        widest = max((len(str(cell.value)) for cell in col if cell.value is not None), default=0)
        sheet.column_dimensions[get_column_letter(col[0].column)].width = widest + 2


def sheet_rows_for(aggregated_rows):
    return [
        [r["instance"], r["initial_cost"], r["final_cost_mean"], r["rpd"], r["rpd_source"], r["time_ms_mean"],
         r["time_ms_min"], r["time_ms_max"], r["time_limit_ms"], r["budget_used_pct"], r["n_iterations"], r["note"]]
        for r in aggregated_rows
    ]


def write_workbook(results, literature, out_path):
    wb = Workbook()
    wb.remove(wb.active)

    # entrada -> linhas agregadas por instância (uma por instância, médias sobre as iterações).
    aggregated_by_entry = {entry: aggregate_by_instance(rows) for entry, rows in results.items()}
    attach_rpd(aggregated_by_entry, literature)

    summary = []
    for entry, agg_rows in aggregated_by_entry.items():
        if not agg_rows:
            continue
        rpds = [r["rpd"] for r in agg_rows]
        times = [r["time_ms_mean"] for r in agg_rows]
        budget_used = [r["budget_used_pct"] for r in agg_rows]
        summary.append((
            entry, sum(rpds) / len(rpds), min(rpds), max(rpds), sum(budget_used) / len(budget_used),
            sum(times) / len(times), len(agg_rows),
        ))
    # Menor RPD médio primeiro: a aba "Resumo" já funciona como ranking.
    summary.sort(key=lambda row: row[1])

    summary_sheet = wb.create_sheet(title="Resumo")
    summary_sheet.append([
        "Entrada", "RPD médio (%)", "RPD mín (%)", "RPD máx (%)", "Uso médio do orçamento (%)",
        "Tempo médio (ms)", "N instâncias",
    ])
    for row in summary:
        summary_sheet.append(list(row))
    autosize_columns(summary_sheet)

    for entry, agg_rows in aggregated_by_entry.items():
        sheet = wb.create_sheet(title=entry[:31])
        sheet.append(SHEET_HEADER)
        for row in sheet_rows_for(agg_rows):
            sheet.append(row)
        autosize_columns(sheet)

    wb.save(out_path)
    return summary


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    script_dir = os.path.dirname(os.path.abspath(__file__))
    parser.add_argument("results", nargs="?", default=os.path.join(script_dir, "results"),
                        help="diretório com os results/*.csv (default: ./results)")
    parser.add_argument("--output", default=os.path.join(script_dir, "output.xlsx"),
                        help="xlsx a gravar (default: ./output.xlsx)")
    parser.add_argument("--literature", default=os.path.join(script_dir, "..", "best_ribas.json"),
                        help="JSON com os melhores custos conhecidos (Taillard/Ribas), usado como referência do "
                             "RPD quando disponível (default: ../best_ribas.json)")
    args = parser.parse_args()

    results = discover_results(args.results)
    literature = load_literature_best(args.literature)
    if not literature:
        print(f"aviso: {args.literature} não encontrado - RPD vai usar o melhor desta rodada em toda instância",
              file=sys.stderr)
    summary = write_workbook(results, literature, args.output)

    total_raw_rows = sum(len(rows) for rows in results.values())
    print(f"{'#':>3} {'Entrada':<26} {'RPD médio':>12} {'mín':>10} {'máx':>10} {'Uso orçamento':>15} "
          f"{'Tempo médio (ms)':>18} {'N inst':>8}")
    for rank, (entry, avg, pmin, pmax, avg_budget, avg_time, count) in enumerate(summary, start=1):
        print(f"{rank:>3} {entry:<26} {avg:>11.3f}% {pmin:>9.3f}% {pmax:>9.3f}% {avg_budget:>14.1f}% "
              f"{avg_time:>18.3f} {count:>8}")

    print(f"\naggregated {total_raw_rows} raw iteration rows across {len(results)} sheets to {args.output}")


if __name__ == "__main__":
    main()
