#!/usr/bin/env python3
"""Agrega os results/<entrada>.csv gerados pelo run_all.py (uma linha bruta
por REPETIÇÃO de um IG_IJ inteiro rodado com aquela busca local no lugar do
passo de busca local - ver README.md) num único output.xlsx com os valores
médios por instância, incluindo RPD e métricas de convergência.

Formato de results/<entrada>.csv (cabeçalho + `iters` linhas por instância):
    instance,iteration,initial_cost,final_best_cost,n_calls,call_of_best,avg_improvement_per_call,time_ms,time_limit_ms,note
    J20M5N1,1,1457,1374,6717,790,30.6107,200.015,200,
    J20M5N1,2,1457,1374,6834,790,30.6405,200.02,200,
    ...

Colunas específicas deste harness (não existem em local-searchs/local-searchs-fair):
- n_calls: quantas vezes o IG chamou a busca local dentro do orçamento de tempo.
- call_of_best: em qual chamada o `best` foi encontrado pela última vez - quanto
  menor em relação a n_calls, mais rápido essa busca local fez o IG convergir
  (e mais "desperdiçado" ficou o resto do orçamento).
- avg_improvement_per_call: melhoria média de custo (antes/depois da busca
  local) a cada chamada - a métrica "melhoria a cada vez que o algoritmo
  chama a local search" pedida na proposta original.

Este script agrupa as linhas brutas por instância e calcula, por entrada, uma
aba com uma linha por instância:

    Instância | Custo inicial | Custo final (méd) | RPD (%) | Fonte RPD
             | N chamadas (méd) | Chamada do melhor (méd) | % do sweep até o melhor
             | Melhoria méd/chamada | Tempo médio (ms) | Orçamento (ms) | N iterações | Nota

Uma aba "Resumo" no início lista, por entrada, o RPD médio/mín/máx e as
métricas de convergência agregadas - ordenada do menor (melhor) para o maior
RPD médio, ou seja, já funciona como um ranking de "qual busca local faz o
IG_IJ chegar mais perto do melhor resultado".

As trajetórias brutas (trajectories/<entrada>__<instância>.csv, uma linha por
chamada) NÃO são agregadas aqui - são para inspeção/plot manual da curva de
convergência (custo x chamada, ou custo x tempo) de uma entrada específica.
"""

import argparse
import csv
import json
import math
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

# Ordem do catálogo (#1..#23) do README, usada nas abas.
ENTRY_ORDER = [
    "rls",
    "rls-grabowski",
    "rls-de-abc",
    "mrls-p-eda",
    "best-swap-ig",
    "best-swap-ig-fixed",
    "swap-first-ig",
    "vnd1",
    "vnd2",
    "vnd1-fixed",
    "vnd2-fixed",
    "ig-ij-dispatch",
    "svns-s-swap",
    "svns-s-insertion",
    "svns-d-swap",
    "svns-d-insertion",
    "hvns-best-insertion",
    "hvns-best-edge-insertion",
    "hvns-best-swap",
    "hvns-shaking",
    "hvns-sa-rls",
    "hvns-sa-edge-insertion",
    "tpa-sa",
]

SHEET_HEADER = [
    "Instância", "Custo inicial", "Custo final (méd)", "RPD (%)", "Fonte RPD", "N chamadas (méd)",
    "Chamada do melhor (méd)", "% até o melhor", "Melhoria méd/chamada", "Tempo médio (ms)", "Orçamento (ms)",
    "N iterações", "Nota",
]

INSTANCE_RE = re.compile(r"^(J\d+M\d+)N(\d+)$")


def load_literature_best(path):
    if not os.path.exists(path):
        return {}
    with open(path) as f:
        return json.load(f)


def literature_best_cost(instance_name, literature):
    m = INSTANCE_RE.match(instance_name)
    if not m:
        return None
    group, index = m.group(1), int(m.group(2)) - 1
    values = literature.get(group)
    if values is None or index >= len(values):
        return None
    return values[index]


def parse_results_file(path):
    rows = []
    with open(path, newline="") as f:
        reader = csv.DictReader(f)
        for rec in reader:
            rows.append({
                "instance": rec["instance"],
                "iteration": int(rec["iteration"]),
                "initial_cost": int(rec["initial_cost"]),
                "final_best_cost": int(rec["final_best_cost"]),
                "n_calls": int(rec["n_calls"]),
                "call_of_best": int(rec["call_of_best"]),
                "avg_improvement_per_call": float(rec["avg_improvement_per_call"]),
                "time_ms": float(rec["time_ms"]),
                "time_limit_ms": float(rec["time_limit_ms"]),
                "note": rec["note"],
            })
    return rows


def aggregate_by_instance(rows):
    order = []
    by_instance = {}
    for r in rows:
        key = r["instance"]
        if key not in by_instance:
            order.append(key)
            by_instance[key] = []
        by_instance[key].append(r)

    aggregated = []
    for instance in order:
        group = by_instance[instance]
        n_calls = [r["n_calls"] for r in group]
        call_of_best = [r["call_of_best"] for r in group]
        improvements = [r["avg_improvement_per_call"] for r in group]
        times = [r["time_ms"] for r in group]
        n_calls_mean = sum(n_calls) / len(n_calls)
        call_of_best_mean = sum(call_of_best) / len(call_of_best)
        aggregated.append({
            "instance": instance,
            "initial_cost": group[0]["initial_cost"],
            "final_best_cost_mean": sum(r["final_best_cost"] for r in group) / len(group),
            "n_calls_mean": n_calls_mean,
            "call_of_best_mean": call_of_best_mean,
            "pct_calls_to_best": (call_of_best_mean / n_calls_mean * 100) if n_calls_mean > 0 else 0.0,
            "avg_improvement_per_call_mean": sum(improvements) / len(improvements),
            "time_ms_mean": sum(times) / len(times),
            "time_limit_ms": group[0]["time_limit_ms"],
            "n_iterations": len(group),
            "note": group[-1]["note"],
        })
    return aggregated


def attach_rpd(aggregated_by_entry, literature):
    best_found = {}
    for agg_rows in aggregated_by_entry.values():
        for r in agg_rows:
            best_found[r["instance"]] = min(best_found.get(r["instance"], math.inf), r["final_best_cost_mean"])

    for agg_rows in aggregated_by_entry.values():
        for r in agg_rows:
            lit_best = literature_best_cost(r["instance"], literature)
            if lit_best:
                reference, source = lit_best, "literatura (Ribas)"
            else:
                reference, source = best_found[r["instance"]], "melhor desta rodada"
            r["rpd"] = ((r["final_best_cost_mean"] - reference) / reference * 100) if reference > 0 else 0.0
            r["rpd_source"] = source


def discover_results(results_dir):
    found = {}
    for name in os.listdir(results_dir):
        if not name.endswith(".csv"):
            continue
        entry = os.path.splitext(name)[0]
        found[entry] = parse_results_file(os.path.join(results_dir, name))
    if not found:
        sys.exit(f"nenhum results/*.csv encontrado em {results_dir} (rode run_all.py antes)")

    ordered = [e for e in ENTRY_ORDER if e in found]
    ordered += sorted(e for e in found if e not in ENTRY_ORDER)
    return {e: found[e] for e in ordered}


def autosize_columns(sheet):
    for col in sheet.iter_cols():
        widest = max((len(str(cell.value)) for cell in col if cell.value is not None), default=0)
        sheet.column_dimensions[get_column_letter(col[0].column)].width = widest + 2


def sheet_rows_for(aggregated_rows):
    return [
        [r["instance"], r["initial_cost"], r["final_best_cost_mean"], r["rpd"], r["rpd_source"], r["n_calls_mean"],
         r["call_of_best_mean"], r["pct_calls_to_best"], r["avg_improvement_per_call_mean"], r["time_ms_mean"],
         r["time_limit_ms"], r["n_iterations"], r["note"]]
        for r in aggregated_rows
    ]


def write_workbook(results, literature, out_path):
    wb = Workbook()
    wb.remove(wb.active)

    aggregated_by_entry = {entry: aggregate_by_instance(rows) for entry, rows in results.items()}
    attach_rpd(aggregated_by_entry, literature)

    summary = []
    for entry, agg_rows in aggregated_by_entry.items():
        if not agg_rows:
            continue
        rpds = [r["rpd"] for r in agg_rows]
        n_calls = [r["n_calls_mean"] for r in agg_rows]
        pct_to_best = [r["pct_calls_to_best"] for r in agg_rows]
        improvements = [r["avg_improvement_per_call_mean"] for r in agg_rows]
        summary.append((
            entry, sum(rpds) / len(rpds), min(rpds), max(rpds), sum(n_calls) / len(n_calls),
            sum(pct_to_best) / len(pct_to_best), sum(improvements) / len(improvements), len(agg_rows),
        ))
    # Menor RPD médio primeiro: a aba "Resumo" já funciona como ranking.
    summary.sort(key=lambda row: row[1])

    summary_sheet = wb.create_sheet(title="Resumo")
    summary_sheet.append([
        "Entrada", "RPD médio (%)", "RPD mín (%)", "RPD máx (%)", "N chamadas (méd)", "% até o melhor (méd)",
        "Melhoria méd/chamada", "N instâncias",
    ])
    for row in summary:
        summary_sheet.append(list(row))
    autosize_columns(summary_sheet)

    for entry, agg_rows in aggregated_by_entry.items():
        sheet = wb.create_sheet(title=entry[:31])
        sheet.append(SHEET_HEADER)
        for row in sheet_rows_for(agg_rows):
            sheet.append(row)
        autosize_columns(sheet)

    wb.save(out_path)
    return summary


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    script_dir = os.path.dirname(os.path.abspath(__file__))
    parser.add_argument("results", nargs="?", default=os.path.join(script_dir, "results"),
                        help="diretório com os results/*.csv (default: ./results)")
    parser.add_argument("--output", default=os.path.join(script_dir, "output.xlsx"),
                        help="xlsx a gravar (default: ./output.xlsx)")
    parser.add_argument("--literature", default=os.path.join(script_dir, "..", "best_ribas.json"),
                        help="JSON com os melhores custos conhecidos (Taillard/Ribas), usado como referência do "
                             "RPD quando disponível (default: ../best_ribas.json)")
    args = parser.parse_args()

    results = discover_results(args.results)
    literature = load_literature_best(args.literature)
    if not literature:
        print(f"aviso: {args.literature} não encontrado - RPD vai usar o melhor desta rodada em toda instância",
              file=sys.stderr)
    summary = write_workbook(results, literature, args.output)

    total_raw_rows = sum(len(rows) for rows in results.values())
    print(f"{'#':>3} {'Entrada':<26} {'RPD médio':>12} {'mín':>10} {'máx':>10} {'N chamadas':>12} "
          f"{'% até melhor':>13} {'Melhoria/chamada':>18} {'N inst':>8}")
    for rank, (entry, avg, pmin, pmax, avg_calls, avg_pct, avg_improve, count) in enumerate(summary, start=1):
        print(f"{rank:>3} {entry:<26} {avg:>11.3f}% {pmin:>9.3f}% {pmax:>9.3f}% {avg_calls:>12.1f} "
              f"{avg_pct:>12.1f}% {avg_improve:>18.3f} {count:>8}")

    print(f"\naggregated {total_raw_rows} raw repetition rows across {len(results)} sheets to {args.output}")


if __name__ == "__main__":
    main()
