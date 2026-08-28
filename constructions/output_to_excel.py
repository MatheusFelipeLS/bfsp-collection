#!/usr/bin/env python3
"""Converte os arquivos de benchmark em constructions/results/ num único
output.xlsx.

Cada results/<algoritmo>.txt vira uma aba (colunas: Instância, Valor Objetivo,
Tempo (ms), Gap (%)). Uma aba "Resumo" no início agrega o gap por algoritmo
(médio/min/max e nº de instâncias), ordenada por gap médio.

Formato de results/<algoritmo>.txt (pares de linhas, terminado por "-"):
    instances/J20M5/J20M5N1
    1462,0.054045
    instances/J20M5/J20M5N2
    1522,0.058538
    -

O nome do algoritmo vem do nome do arquivo.

gap% = (valor_obtido - melhor_conhecido) / melhor_conhecido * 100
       (melhor conhecido: best_ribas.json; instâncias sem entrada ficam com gap
       em branco e são ignoradas no Resumo)
"""

import argparse
import glob
import json
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

DATA_LINE_RE = re.compile(r"^(\d+),([0-9.eE+-]+)$")
INSTANCE_RE = re.compile(r"^(?P<group>[A-Za-z0-9]+)N(?P<n>\d+)$")


def parse_results_file(path):
    """Retorna [(instance, cost, time_ms), ...] de um results/<algoritmo>.txt."""

    rows = []
    current_instance = None

    with open(path) as f:
        for lineno, raw_line in enumerate(f, start=1):
            line = raw_line.strip()

            if not line or line == "-":
                continue

            match = DATA_LINE_RE.match(line)
            if match:
                if current_instance is None:
                    raise ValueError(f"{path}:{lineno}: linha de resultado antes de qualquer instância: {line!r}")
                cost, time_ms = match.groups()
                rows.append((current_instance, int(cost), float(time_ms)))
                continue

            if "/" in line:
                current_instance = os.path.basename(line)
            else:
                print(f"warning: {path}:{lineno}: linha ignorada: {line!r}", file=sys.stderr)

    return rows


def parse_results_dir(path):
    """Retorna {algoritmo: [(instance, cost, time_ms), ...]}, algoritmos em ordem
    alfabética case-insensitive (mesma ordem do run_all.py)."""

    files = sorted(glob.glob(os.path.join(path, "*.txt")), key=lambda p: os.path.basename(p).lower())
    if not files:
        sys.exit(f"nenhum results/*.txt encontrado em {path}")

    return {os.path.splitext(os.path.basename(f))[0]: parse_results_file(f) for f in files}


def load_best(path):
    with open(path) as f:
        return json.load(f)


def best_value(best, instance):
    m = INSTANCE_RE.match(instance)
    if not m:
        return None
    values = best.get(m.group("group"))
    n = int(m.group("n"))
    if values is None or n < 1 or n > len(values):
        return None
    return values[n - 1]


def autosize_columns(sheet):
    """Ajusta a largura de cada coluna ao maior valor (cabeçalho incluído)."""
    for col in sheet.iter_cols():
        widest = max((len(str(cell.value)) for cell in col if cell.value is not None), default=0)
        sheet.column_dimensions[get_column_letter(col[0].column)].width = widest + 2


def write_workbook(results, best, out_path):
    """Escreve o workbook e retorna (summary, skipped_groups)."""

    wb = Workbook()
    wb.remove(wb.active)

    summary_sheet = wb.create_sheet(title="Resumo")
    summary_sheet.append(["Algoritmo", "Gap médio %", "Gap min %", "Gap max %", "N instâncias"])

    skipped_groups = set()
    summary = []
    algo_sheets = []

    for algorithm, rows in results.items():
        sheet_rows = []
        gaps = []
        for instance, cost, time_ms in rows:
            bv = best_value(best, instance)
            if bv is None:
                m = INSTANCE_RE.match(instance)
                skipped_groups.add(m.group("group") if m else instance)
                gap = None
            else:
                gap = (cost - bv) / bv * 100
                gaps.append(gap)
            sheet_rows.append([instance, cost, time_ms, gap])

        algo_sheets.append((algorithm, sheet_rows))
        if gaps:
            summary.append((algorithm, sum(gaps) / len(gaps), min(gaps), max(gaps), len(gaps)))

    summary.sort(key=lambda row: row[1])
    for row in summary:
        summary_sheet.append(list(row))
    autosize_columns(summary_sheet)

    for algorithm, sheet_rows in algo_sheets:
        sheet = wb.create_sheet(title=algorithm[:31])
        sheet.append(["Instância", "Valor Objetivo", "Tempo (ms)", "Gap (%)"])
        for row in sheet_rows:
            sheet.append(row)
        autosize_columns(sheet)

    wb.save(out_path)
    return summary, skipped_groups


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    script_dir = os.path.dirname(os.path.abspath(__file__))
    parser.add_argument("results", nargs="?", default=os.path.join(script_dir, "results"),
                        help="diretório com os results/*.txt (default: ./results)")
    parser.add_argument("--best", default=os.path.join(os.path.dirname(script_dir), "best_ribas.json"),
                        help="JSON com os melhores valores conhecidos (default: ../best_ribas.json)")
    parser.add_argument("--output", default=os.path.join(script_dir, "output.xlsx"),
                        help="xlsx a gravar (default: ./output.xlsx)")
    args = parser.parse_args()

    results = parse_results_dir(args.results)
    best = load_best(args.best)

    summary, skipped_groups = write_workbook(results, best, args.output)

    total_rows = sum(len(rows) for rows in results.values())
    print(f"{'Algoritmo':<12} {'Gap médio %':>12} {'Gap min %':>12} {'Gap max %':>12} {'N instâncias':>14}")
    for algorithm, avg, gmin, gmax, count in summary:
        print(f"{algorithm:<12} {avg:>12.3f} {gmin:>12.3f} {gmax:>12.3f} {count:>14}")

    print(f"\nwrote {total_rows} results across {len(results)} sheets to {args.output}")

    if skipped_groups:
        print(
            f"(grupos sem valor em best_ribas.json, gap em branco: {', '.join(sorted(skipped_groups))})",
            file=sys.stderr,
        )


if __name__ == "__main__":
    main()
