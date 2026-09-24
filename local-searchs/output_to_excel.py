#!/usr/bin/env python3
"""Agrega os results/<entrada>.csv gerados pelo run_all.py (uma linha bruta
por ITERAÇÃO) num único output.xlsx com os valores médios por instância.

Formato de results/<entrada>.csv (cabeçalho + `iters` linhas por instância,
uma por repetição - sem nenhuma agregação prévia):
    instance,iteration,initial_cost,final_cost,improvement_abs,improvement_pct,time_ms,note
    J20M5N1,1,1828,1447,381,20.8425,0.061201,
    J20M5N1,2,1828,1447,381,20.8425,0.058863,
    ...

Este script agrupa essas linhas por instância e calcula, por entrada, uma
aba com uma linha por instância:

    Instância | Custo inicial | Custo final (méd) | Melhoria absoluta (méd)
             | Melhoria (%) (méd) | Tempo médio (ms) | Tempo mín (ms)
             | Tempo máx (ms) | N iterações | Nota

Uma aba "Resumo" no início agrega, por entrada, a melhoria média/mín/máx (%)
e o tempo médio, já em cima dos valores por-instância acima descritos.
"""

import argparse
import csv
import os
import sys

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

# Ordem do catálogo (#1..#23) do README, usada nas abas e no Resumo.
ENTRY_ORDER = [
    "rls",
    "rls-grabowski",
    "rls-de-abc",
    "mrls-p-eda",
    "best-swap-ig",
    "best-swap-ig-fixed",
    "swap-first-ig",
    "vnd1",
    "vnd2",
    "vnd1-fixed",
    "vnd2-fixed",
    "ig-ij-dispatch",
    "svns-s-swap",
    "svns-s-insertion",
    "svns-d-swap",
    "svns-d-insertion",
    "hvns-best-insertion",
    "hvns-best-edge-insertion",
    "hvns-best-swap",
    "hvns-shaking",
    "hvns-sa-rls",
    "hvns-sa-edge-insertion",
    "tpa-sa",
]

SHEET_HEADER = [
    "Instância", "Custo inicial", "Custo final (méd)", "Melhoria absoluta (méd)", "Melhoria (%) (méd)",
    "Tempo médio (ms)", "Tempo mín (ms)", "Tempo máx (ms)", "N iterações", "Nota",
]


def parse_results_file(path):
    """Retorna [{instance, iteration, initial_cost, final_cost, improvement_abs,
    improvement_pct, time_ms, note}, ...] - uma entrada por iteração bruta."""
    rows = []
    with open(path, newline="") as f:
        reader = csv.DictReader(f)
        for rec in reader:
            rows.append({
                "instance": rec["instance"],
                "iteration": int(rec["iteration"]),
                "initial_cost": int(rec["initial_cost"]),
                "final_cost": int(rec["final_cost"]),
                "improvement_abs": float(rec["improvement_abs"]),
                "improvement_pct": float(rec["improvement_pct"]),
                "time_ms": float(rec["time_ms"]),
                "note": rec["note"],
            })
    return rows


def aggregate_by_instance(rows):
    """Agrupa as linhas brutas (uma por iteração) por instância e retorna uma
    linha agregada por instância, na ordem de primeira aparição."""
    order = []
    by_instance = {}
    for r in rows:
        key = r["instance"]
        if key not in by_instance:
            order.append(key)
            by_instance[key] = []
        by_instance[key].append(r)

    aggregated = []
    for instance in order:
        group = by_instance[instance]
        times = [r["time_ms"] for r in group]
        aggregated.append({
            "instance": instance,
            "initial_cost": group[0]["initial_cost"],
            "final_cost_mean": sum(r["final_cost"] for r in group) / len(group),
            "improvement_abs_mean": sum(r["improvement_abs"] for r in group) / len(group),
            "improvement_pct_mean": sum(r["improvement_pct"] for r in group) / len(group),
            "time_ms_mean": sum(times) / len(times),
            "time_ms_min": min(times),
            "time_ms_max": max(times),
            "n_iterations": len(group),
            "note": group[-1]["note"],
        })
    return aggregated


def discover_results(results_dir):
    """Retorna {entrada: rows_brutas} (uma linha por iteração) para os
    results/<entrada>.csv existentes, na ordem de ENTRY_ORDER (extras vão
    para o fim, em ordem alfabética)."""
    found = {}
    for name in os.listdir(results_dir):
        if not name.endswith(".csv"):
            continue
        entry = os.path.splitext(name)[0]
        found[entry] = parse_results_file(os.path.join(results_dir, name))
    if not found:
        sys.exit(f"nenhum results/*.csv encontrado em {results_dir} (rode run_all.py antes)")

    ordered = [e for e in ENTRY_ORDER if e in found]
    ordered += sorted(e for e in found if e not in ENTRY_ORDER)
    return {e: found[e] for e in ordered}


def autosize_columns(sheet):
    """Ajusta a largura de cada coluna ao maior valor (cabeçalho incluído)."""
    for col in sheet.iter_cols():
        widest = max((len(str(cell.value)) for cell in col if cell.value is not None), default=0)
        sheet.column_dimensions[get_column_letter(col[0].column)].width = widest + 2


def sheet_rows_for(aggregated_rows):
    return [
        [r["instance"], r["initial_cost"], r["final_cost_mean"], r["improvement_abs_mean"],
         r["improvement_pct_mean"], r["time_ms_mean"], r["time_ms_min"], r["time_ms_max"], r["n_iterations"],
         r["note"]]
        for r in aggregated_rows
    ]


def write_workbook(results, out_path):
    wb = Workbook()
    wb.remove(wb.active)

    # entrada -> linhas agregadas por instância (uma por instância, médias sobre as iterações).
    aggregated_by_entry = {entry: aggregate_by_instance(rows) for entry, rows in results.items()}

    summary_sheet = wb.create_sheet(title="Resumo")
    summary_sheet.append(
        ["Entrada", "Melhoria média (%)", "Melhoria mín (%)", "Melhoria máx (%)", "Tempo médio (ms)", "N instâncias"]
    )

    summary = []
    for entry, agg_rows in aggregated_by_entry.items():
        if not agg_rows:
            continue
        pcts = [r["improvement_pct_mean"] for r in agg_rows]
        times = [r["time_ms_mean"] for r in agg_rows]
        summary.append(
            (entry, sum(pcts) / len(pcts), min(pcts), max(pcts), sum(times) / len(times), len(agg_rows))
        )

    for row in summary:
        summary_sheet.append(list(row))
    autosize_columns(summary_sheet)

    for entry, agg_rows in aggregated_by_entry.items():
        sheet = wb.create_sheet(title=entry[:31])
        sheet.append(SHEET_HEADER)
        for row in sheet_rows_for(agg_rows):
            sheet.append(row)
        autosize_columns(sheet)

    wb.save(out_path)
    return summary


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    script_dir = os.path.dirname(os.path.abspath(__file__))
    parser.add_argument("results", nargs="?", default=os.path.join(script_dir, "results"),
                        help="diretório com os results/*.csv (default: ./results)")
    parser.add_argument("--output", default=os.path.join(script_dir, "output.xlsx"),
                        help="xlsx a gravar (default: ./output.xlsx)")
    args = parser.parse_args()

    results = discover_results(args.results)
    summary = write_workbook(results, args.output)

    total_raw_rows = sum(len(rows) for rows in results.values())
    print(f"{'Entrada':<26} {'Melhoria média':>15} {'mín':>10} {'máx':>10} {'Tempo médio (ms)':>18} {'N inst':>8}")
    for entry, avg, pmin, pmax, avg_time, count in summary:
        print(f"{entry:<26} {avg:>14.3f}% {pmin:>9.3f}% {pmax:>9.3f}% {avg_time:>18.3f} {count:>8}")

    print(f"\naggregated {total_raw_rows} raw iteration rows across {len(results)} sheets to {args.output}")


if __name__ == "__main__":
    main()
