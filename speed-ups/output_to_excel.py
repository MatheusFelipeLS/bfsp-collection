#!/usr/bin/env python3
"""Converte os results/<speedup>.csv gerados pelo run_all.py num único
output.xlsx.

Uma aba "Resumo" no início agrega o fator de aceleração por speed up
(médio/min/max, nº de instâncias e nº de instâncias com resultado verificado).
Depois, uma aba por speed up, na ordem #1..#6, com só as colunas que fazem
sentido para aquele speed up:

  - taillard-insertion, partial-recalc, edge-insertion, alpha-sigma,
    preallocated-buffers  -> resultado idêntico ao ingênuo, só muda o tempo:
        Instância | Tempo ingênuo (ms) | Tempo acelerado (ms) | Fator | Correto
  - grabowski-block  -> compara duas buscas locais diferentes (restrita vs.
    completa), então tempo E qualidade importam:
        Instância | Tempo rls (ms) | Tempo grabowski (ms) | Fator
                  | Custo inicial | Custo rls | Custo grabowski | Correto

Formato de results/<speedup>.csv (cabeçalho + uma linha por instância):
    instance,naive_ms,accel_ms,factor,correct,note
    J20M5N1,0.248272,0.0153823,16.1401,1,makespan=1463
"""

import argparse
import csv
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

# Ordem natural dos speed ups (#1..#6), usada nas abas e no Resumo.
SPEEDUP_ORDER = [
    "taillard-insertion",
    "grabowski-block",
    "partial-recalc",
    "edge-insertion",
    "alpha-sigma",
    "preallocated-buffers",
]

TIME_ONLY_HEADER = ["Instância", "Tempo ingênuo (ms)", "Tempo acelerado (ms)", "Fator", "Correto"]
GRABOWSKI_HEADER = [
    "Instância", "Tempo rls (ms)", "Tempo grabowski (ms)", "Fator",
    "Custo inicial", "Custo rls", "Custo grabowski", "Correto",
]

GRABOWSKI_NOTE_RE = re.compile(r"start=(\d+)\s+restricted=(\d+)\s+full=(\d+)")


def parse_results_file(path):
    """Retorna [{instance, naive_ms, accel_ms, factor, correct, note}, ...]."""
    rows = []
    with open(path, newline="") as f:
        reader = csv.DictReader(f)
        for rec in reader:
            rows.append({
                "instance": rec["instance"],
                "naive_ms": float(rec["naive_ms"]),
                "accel_ms": float(rec["accel_ms"]),
                "factor": float(rec["factor"]),
                "correct": int(rec["correct"]),
                "note": rec["note"],
            })
    return rows


def discover_results(results_dir):
    """Retorna {speedup: rows} para os results/<speedup>.csv existentes,
    na ordem de SPEEDUP_ORDER (extras vão para o fim, em ordem alfabética)."""
    found = {}
    for name in os.listdir(results_dir):
        if not name.endswith(".csv"):
            continue
        speedup = os.path.splitext(name)[0]
        found[speedup] = parse_results_file(os.path.join(results_dir, name))
    if not found:
        sys.exit(f"nenhum results/*.csv encontrado em {results_dir} (rode run_all.py antes)")

    ordered = [s for s in SPEEDUP_ORDER if s in found]
    ordered += sorted(s for s in found if s not in SPEEDUP_ORDER)
    return {s: found[s] for s in ordered}


def autosize_columns(sheet):
    """Ajusta a largura de cada coluna ao maior valor (cabeçalho incluído)."""
    for col in sheet.iter_cols():
        widest = max((len(str(cell.value)) for cell in col if cell.value is not None), default=0)
        sheet.column_dimensions[get_column_letter(col[0].column)].width = widest + 2


def sheet_rows_for(speedup, rows):
    """Converte os registros brutos nas linhas da aba daquele speed up."""
    out = []
    for r in rows:
        if speedup == "grabowski-block":
            m = GRABOWSKI_NOTE_RE.search(r["note"])
            start, restricted, full = (int(x) for x in m.groups()) if m else (None, None, None)
            out.append([
                r["instance"], r["naive_ms"], r["accel_ms"], r["factor"],
                start, restricted, full, r["correct"],
            ])
        else:
            out.append([r["instance"], r["naive_ms"], r["accel_ms"], r["factor"], r["correct"]])
    return out


def write_workbook(results, out_path):
    wb = Workbook()
    wb.remove(wb.active)

    summary_sheet = wb.create_sheet(title="Resumo")
    summary_sheet.append(["Speedup", "Fator médio", "Fator mínimo", "Fator máximo", "N instâncias", "N corretos"])

    summary = []
    for speedup, rows in results.items():
        if not rows:
            continue
        factors = [r["factor"] for r in rows]
        n_correct = sum(r["correct"] for r in rows)
        summary.append((speedup, sum(factors) / len(factors), min(factors), max(factors), len(rows), n_correct))

    for row in summary:
        summary_sheet.append(list(row))
    autosize_columns(summary_sheet)

    for speedup, rows in results.items():
        sheet = wb.create_sheet(title=speedup[:31])
        header = GRABOWSKI_HEADER if speedup == "grabowski-block" else TIME_ONLY_HEADER
        sheet.append(header)
        for row in sheet_rows_for(speedup, rows):
            sheet.append(row)
        autosize_columns(sheet)

    wb.save(out_path)
    return summary


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    script_dir = os.path.dirname(os.path.abspath(__file__))
    parser.add_argument("results", nargs="?", default=os.path.join(script_dir, "results"),
                        help="diretório com os results/*.csv (default: ./results)")
    parser.add_argument("--output", default=os.path.join(script_dir, "output.xlsx"),
                        help="xlsx a gravar (default: ./output.xlsx)")
    args = parser.parse_args()

    results = discover_results(args.results)
    summary = write_workbook(results, args.output)

    total_rows = sum(len(rows) for rows in results.values())
    print(f"{'Speedup':<22} {'Fator médio':>12} {'Fator min':>12} {'Fator max':>12} {'N inst':>8} {'N ok':>6}")
    for speedup, avg, fmin, fmax, count, n_ok in summary:
        print(f"{speedup:<22} {avg:>12.3f} {fmin:>12.3f} {fmax:>12.3f} {count:>8} {n_ok:>6}")

    print(f"\nwrote {total_rows} results across {len(results)} sheets to {args.output}")


if __name__ == "__main__":
    main()
